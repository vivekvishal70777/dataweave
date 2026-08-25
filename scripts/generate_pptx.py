#!/usr/bin/env python3
"""Build simple 16:9 PowerPoint decks for screen-share recording."""
from __future__ import annotations

import re
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import generate_slides as gs  # noqa: E402

OUT = ROOT / "instructor" / "slides" / "pptx"
ART = ROOT / "instructor" / "slides" / "assets" / "art"

NAVY = RGBColor(0x0F, 0x27, 0x44)
AMBER_BG = RGBColor(0x3D, 0x29, 0x10)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xB7, 0xCB, 0xDE)
ACCENT = RGBColor(0x3E, 0xC6, 0xFF)
AMBER = RGBColor(0xFF, 0xB0, 0x20)
CODE_BG = RGBColor(0x06, 0x12, 0x1C)
CODE_FG = RGBColor(0xD7, 0xE8, 0xF5)

W = Inches(13.333)
H = Inches(7.5)

# First matching keyword wins.
ART_RULES: list[tuple[tuple[str, ...], str]] = [
    (("pause",), "dw-pause.jpg"),
    (("index.pptx", "recording index"), "dw-hero-transform.jpg"),
    (("s14", "mapping", "canonical", "cluster", "14-dataweave"), "dw-mapping.jpg"),
    (("interview", "whiteboard", "bootcamp", "l54", "l55", "l56", "l58", "question-bank"), "dw-interview.jpg"),
    (("stream", "crypto", "hash", "hmac", "performance", "l50", "l51", "l52"), "dw-streaming.jpg"),
    (("recurs", "pii", "namespace", "mask", "l44", "l45", "l46", "l47", "l48", "diff", "dynamic-key"), "dw-recursion.jpg"),
    (("leftjoin", "left-join", "lookup", "l39", "l40", "l41", "l42", "attribute", "mule-context", "joins-modules"), "dw-joins.jpg"),
    (("date", "match", "try", "l35", "l36", "l37", "timezone"), "dw-dates.jpg"),
    (("group", "reduce", "flatmap", "merge", "l29", "l30", "l31", "l32", "l33", "orderby"), "dw-group-reduce.jpg"),
    (("xml", "csv", "json-xml", "l24", "l25", "l26", "l27", "format", "mime"), "dw-formats.jpg"),
    (("string", "if-else", "if/else", "plus-plus", "l20", "l21", "l22", "split", "join sku"), "dw-strings.jpg"),
    (("object", "null", "pluck", "selector", "mapobject", "l15", "l16", "l17", "l18"), "dw-objects.jpg"),
    (("map-filter", "filter", "sizeof", "flatten", "l11", "l12", "l13", "array"), "dw-map-filter.jpg"),
    (("script", "var", "fun", "type", "l05", "l06", "l07", "l08", "dataweave"), "dw-script.jpg"),
    (("folder", "l04", "student-folder"), "dw-folders.jpg"),
    (("playground", "l03", "getting-started"), "dw-playground.jpg"),
    (("loop", "l02", "how this course", "how-this-course"), "dw-loop.jpg"),
    (("money", "gst", "tax", "invoice", "fx", "currency"), "dw-money.jpg"),
    (("welcome", "l01"), "dw-hero-transform.jpg"),
    (("lab",), "dw-lab.jpg"),
]


def strip_md(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return re.sub(r"\s+", " ", text).strip()


def pick_art(path: Path, deck: gs.Deck, item: dict) -> Path | None:
    if item.get("kind") == "pause":
        return ART / "dw-pause.jpg"
    heading = (item.get("heading") or "").lower()
    if heading == "the loop":
        return ART / "dw-loop.jpg"
    blob = " ".join(
        [
            path.as_posix().lower(),
            path.stem.lower(),
            deck.kicker.lower(),
            deck.title.lower(),
            heading,
            item.get("kind", ""),
        ]
    )
    for keys, name in ART_RULES:
        if any(k in blob for k in keys):
            p = ART / name
            return p if p.exists() else None
    fallback = ART / "dw-hero-transform.jpg"
    return fallback if fallback.exists() else None


def set_run(run, size: int, color: RGBColor, bold: bool = False, font: str = "Calibri") -> None:
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font


def fill_shape(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_notes(slide, notes: str) -> None:
    if not notes:
        return
    tf = slide.notes_slide.notes_text_frame
    tf.text = notes


def add_picture(slide, image: Path | None, left: float, top: float, size: float) -> None:
    if not image or not image.exists():
        return
    slide.shapes.add_picture(str(image), Inches(left), Inches(top), Inches(size), Inches(size))


def new_slide(prs: Presentation, bg: RGBColor):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    fill_shape(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H), bg)
    return slide


def kicker_box(slide, text: str, color: RGBColor) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(0.28), Inches(12), Inches(0.4))
    p = box.text_frame.paragraphs[0]
    p.text = text.upper()
    set_run(p.runs[0], 14, color, bold=True)


def title_box(slide, text: str, top: float, width: float, size: int = 36) -> None:
    box = slide.shapes.add_textbox(Inches(0.7), Inches(top), Inches(width), Inches(1.8))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = strip_md(text)
    set_run(p.runs[0], size, WHITE, bold=True)


def bullets_box(slide, bullets: list[str], top: float, width: float, color: RGBColor = WHITE) -> None:
    box = slide.shapes.add_textbox(Inches(0.85), Inches(top), Inches(width), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "•  " + strip_md(item)
        p.space_after = Pt(14)
        set_run(p.runs[0], 22, color)


def code_box(slide, code: str) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.05), Inches(11.9), Inches(4.7)
    )
    fill_shape(shape, CODE_BG)
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.28)
    tf.margin_top = Inches(0.18)
    lines = code.rstrip().splitlines() or [""]
    if len(lines) > 16:
        lines = lines[:15] + ["…"]
    size = 16 if len(lines) <= 10 else 13 if len(lines) <= 14 else 12
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line if line else " "
        p.space_after = Pt(2)
        set_run(p.runs[0], size, CODE_FG, font="Consolas")


def title_slide(prs: Presentation, item: dict, image: Path | None) -> None:
    slide = new_slide(prs, NAVY)
    kicker_box(slide, item["kicker"], ACCENT)
    title_box(slide, item["heading"], 2.0, 7.3 if image else 12, 40)
    if item.get("subtitle"):
        box = slide.shapes.add_textbox(Inches(0.7), Inches(4.15), Inches(7.3 if image else 11.8), Inches(1.2))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = strip_md(item["subtitle"])
        set_run(p.runs[0], 18, MUTED)
    chips = item.get("chips") or []
    if chips:
        box = slide.shapes.add_textbox(Inches(0.7), Inches(5.6), Inches(7.4 if image else 12), Inches(0.5))
        p = box.text_frame.paragraphs[0]
        p.text = "   ·   ".join(strip_md(c) for c in chips)
        set_run(p.runs[0], 15, ACCENT)
    add_picture(slide, image, 8.15, 1.45, 4.55)
    add_notes(slide, item.get("notes", ""))


def content_slide(prs: Presentation, item: dict, image: Path | None) -> None:
    pause = item["kind"] == "pause"
    slide = new_slide(prs, AMBER_BG if pause else NAVY)
    kicker_box(slide, item["kicker"], AMBER if pause else ACCENT)
    use_art = bool(image) and item["kind"] != "code"
    title_box(slide, item["heading"], 0.72, 8.6 if use_art else 12, 30)
    if item["kind"] == "code" and item.get("code"):
        code_box(slide, item["code"])
    elif item.get("heading") == "The loop":
        bullets_box(
            slide,
            [
                "Concept — one idea, 3–7 minutes",
                "Pause — hold the card, students try",
                "Lab — type from the starter, not the solution",
                "Quiz — Udemy practice test, no video",
            ],
            2.55,
            8.2 if use_art else 11.6,
        )
        add_picture(slide, image, 9.35, 2.15, 3.4)
    elif item.get("bullets"):
        bullets_box(slide, item["bullets"], 2.55, 8.2 if use_art else 11.6, AMBER if pause else WHITE)
        if pause:
            add_picture(slide, image, 9.15, 2.05, 3.55)
        else:
            add_picture(slide, image, 9.55, 2.25, 3.15)
    add_notes(slide, item.get("notes", ""))


def write_pptx(path: Path, deck: gs.Deck) -> None:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    items = deck.items or [
        {
            "kind": "title",
            "heading": deck.title,
            "kicker": deck.kicker,
            "subtitle": "",
            "chips": [],
            "notes": "",
            "bullets": [],
            "code": "",
        }
    ]
    for item in items:
        art = pick_art(path, deck, item)
        if item["kind"] == "title":
            title_slide(prs, item, art)
        else:
            content_slide(prs, item, art)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(path))


def parse_index_html(html: str) -> list[tuple[str, list[tuple[str, str]]]]:
    import html as html_lib

    groups: list[tuple[str, list[tuple[str, str]]]] = []
    for heading, table in re.findall(r"<h2>(.*?)</h2>\s*<table>(.*?)</table>", html, flags=re.S):
        items = [
            (html_lib.unescape(kind.strip()), html_lib.unescape(title.strip()))
            for kind, title in re.findall(
                r'class="kind">(.*?)</td>\s*<td><a href="[^"]+">(.*?)</a>',
                table,
                flags=re.S,
            )
        ]
        if items:
            groups.append((html_lib.unescape(heading.strip()), items))
    return groups


def index_deck() -> gs.Deck:
    src = gs.OUT / "index.html"
    html = src.read_text(encoding="utf-8") if src.exists() else ""
    deck = gs.Deck("Index", "Recording slides", "../assets")
    deck.add_title(
        "Course catalog as a shareable deck. F5, then open the matching lecture or lab PPTX.",
        ["Converted from instructor/slides/index.html", "16:9 screen share"],
    )
    deck.add_bullets(
        "How to use this file",
        [
            "Share this window for a course walkthrough or recording plan.",
            "Each later slide is one section from the HTML index.",
            "When you record a video, switch to that file under pptx/lectures or pptx/labs.",
            "Pause cards stay in the lab decks — not in this index.",
        ],
    )
    groups = parse_index_html(html)
    if not groups:
        deck.add_bullets("Missing index", ["Run python3 scripts/generate_slides.py first."])
        return deck
    agenda = [
        f"{i + 1}. {name} ({len(items)})"
        for i, (name, items) in enumerate(groups)
        if not name.startswith("Tutorials")
    ]
    deck.add_bullets("In this index", agenda)
    for name, items in groups:
        bullets = [f"{kind} — {title}" for kind, title in items]
        deck.add_bullets(name, bullets)
    deck.add_bullets(
        "Next",
        [
            "Open pptx/lectures/L01.pptx and press F5 to start recording.",
            "Rebuild this file with python3 scripts/generate_pptx.py.",
        ],
    )
    return deck


def main() -> None:
    n = 0
    write_pptx(OUT / "index.pptx", index_deck())
    n += 1
    write_pptx(ROOT / "instructor" / "slides" / "index.pptx", index_deck())
    n += 1
    for sid, title in gs.SECTION_FOLDERS:
        lec = gs.SECTIONS / sid / "LECTURE.md"
        if lec.exists():
            write_pptx(OUT / "sections" / f"{sid}.pptx", gs.section_overview_deck(lec))
            n += 1

    for path in sorted((gs.TRANSCRIPTS / "lectures").glob("L*.md")):
        write_pptx(OUT / "lectures" / f"{path.stem}.pptx", gs.lecture_deck(path))
        n += 1

    for dest, deck in gs.mapping_cluster_decks():
        write_pptx(OUT / "lectures" / f"{dest.stem}.pptx", deck)
        n += 1

    for path in sorted((gs.TRANSCRIPTS / "labs").glob("*.md")):
        write_pptx(OUT / "labs" / f"{path.stem}.pptx", gs.lab_deck(path))
        n += 1

    for sid, title in gs.SECTION_FOLDERS:
        tdir = gs.TUTORIALS / sid
        if not tdir.is_dir():
            continue
        for path in sorted(tdir.glob("*.md")):
            if path.name.lower() == "readme.md":
                continue
            write_pptx(
                OUT / "tutorials" / sid / f"{path.stem}.pptx",
                gs.tutorial_deck(path, title),
            )
            n += 1

    print(f"Wrote {n} PowerPoint files under {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
